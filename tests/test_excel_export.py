from datetime import datetime, timezone

from openpyxl import load_workbook

from UK_news_scraper.calendar_utils import CalendarMode
from UK_news_scraper.excel_exporter import (
    ExportOptions,
    _ensure_translations_present,
    _translate_texts,
    _translate_with_deep_translator,
    export_news,
)
from UK_news_scraper.models import NewsItem, ParliamentBriefing


def test_export_contains_required_sheets_and_preserves_distinct_links(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "UK_news_scraper.excel_exporter._translate_texts",
        lambda texts, content_label: {},
    )
    monkeypatch.setattr("UK_news_scraper.excel_exporter._ensure_translations_present", lambda *args: None)
    items = [
        NewsItem("Agency", "Agency", "A", "Same title", link, datetime(2026, 6, 8, tzinfo=timezone.utc))
        for link in ("https://example.com/a", "https://example.com/b")
    ]
    parliament = [
        ParliamentBriefing(
            published_at=datetime(2026, 6, 8, tzinfo=timezone.utc),
            chamber="House of Commons",
            publisher="House of Commons Library",
            title="Research title",
            summary="Summary",
            identifier="CBP-1",
            webpage_url="https://example.com/research",
        )
    ]

    path = export_news(items, [], tmp_path / "report.xlsx", parliament_items=parliament)
    workbook = load_workbook(path, read_only=True, data_only=True)

    assert workbook.sheetnames == ["全部新聞", "已初步篩選工作表", "國會研究資料", "篩選設定"]
    assert workbook["全部新聞"].max_row == 5


def test_export_highlights_only_relevant_content_fields(tmp_path, monkeypatch):
    monkeypatch.setattr("UK_news_scraper.excel_exporter._translate_texts", lambda texts, content_label: {})
    monkeypatch.setattr("UK_news_scraper.excel_exporter._ensure_translations_present", lambda *args: None)
    item = NewsItem(
        "Agency", "Agency", "A", "Artificial intelligence policy", "https://example.com/ai",
        datetime(2026, 6, 8, tzinfo=timezone.utc), summary="General announcement",
        matched_topics=["AI"], matched_keywords=["artificial intelligence"],
        title_matched_keywords=["artificial intelligence"], relevance_score=4, relevance_level="低",
    )

    path = export_news([item], [item], tmp_path / "report.xlsx")
    workbook = load_workbook(path)
    all_sheet = workbook["全部新聞"]
    filtered_sheet = workbook["已初步篩選工作表"]

    assert all_sheet["E2"].fill.fgColor.rgb == "00FFF2CC"
    assert all_sheet["A2"].fill.fill_type is None
    assert filtered_sheet["I3"].value == "低"
    assert filtered_sheet["J3"].value == 4


def test_relevance_fill_uses_darker_yellow_for_higher_relevance(tmp_path, monkeypatch):
    monkeypatch.setattr("UK_news_scraper.excel_exporter._translate_texts", lambda texts, content_label: {})
    monkeypatch.setattr("UK_news_scraper.excel_exporter._ensure_translations_present", lambda *args: None)
    levels = [("低", 4), ("中", 6), ("高", 9)]
    items = [
        NewsItem(
            "Agency", "Agency", "A", f"AI policy {level}", f"https://example.com/{score}",
            datetime(2026, 6, 8, tzinfo=timezone.utc), matched_topics=["AI"],
            matched_keywords=["AI"], title_matched_keywords=["AI"],
            relevance_score=score, relevance_level=level,
        )
        for level, score in levels
    ]

    path = export_news(items, items, tmp_path / "relevance-colours.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["全部新聞"]

    colours_by_title = {
        sheet.cell(row, 5).value: sheet.cell(row, 5).fill.fgColor.rgb
        for row in (2, 4, 6)
    }
    assert colours_by_title == {
        "AI policy 高": "00FFD966",
        "AI policy 中": "00FFE699",
        "AI policy 低": "00FFF2CC",
    }


def test_untranslated_titles_only_warn(monkeypatch, capsys):
    item = NewsItem(
        "NCSC",
        "NCSC",
        "NCSC",
        "Cyber Shield: The path to an agentic AI future for cyber defence",
        "https://example.com",
        datetime(2026, 7, 7, tzinfo=timezone.utc),
    )

    _ensure_translations_present({item.title: item.title}, [item])
    captured = capsys.readouterr()
    assert "新聞標題翻譯失敗或未變更" in captured.out


def test_manual_translation_fallback():
    assert (
        _translate_with_deep_translator("Cyber Shield: The path to an agentic AI future for cyber defence")
        == "Cyber Shield：邁向具代理式 AI 的網路防禦未來"
    )


def test_cached_untranslated_title_is_repaired(monkeypatch):
    title = "Cyber Shield: The path to an agentic AI future for cyber defence"
    monkeypatch.setattr("UK_news_scraper.excel_exporter.load_translations", lambda: {title: title})
    monkeypatch.setattr("UK_news_scraper.excel_exporter._translate_uncached_texts", lambda missing, label: {})
    monkeypatch.setattr("UK_news_scraper.excel_exporter.save_translations", lambda translations: None)

    result = _translate_texts([title], "新聞標題")

    assert result[title] == "Cyber Shield：邁向具代理式 AI 的網路防禦未來"


def test_export_uses_roc_number_format_and_strength_columns(tmp_path, monkeypatch):
    monkeypatch.setattr("UK_news_scraper.excel_exporter._translate_texts", lambda texts, content_label: {})
    monkeypatch.setattr("UK_news_scraper.excel_exporter._ensure_translations_present", lambda *args: None)
    item = NewsItem(
        "Agency",
        "Agency",
        "A",
        "Artificial intelligence policy",
        "https://example.com/ai",
        datetime(2026, 6, 8, tzinfo=timezone.utc),
        matched_topics=["AI"],
        matched_keywords=["artificial intelligence"],
        title_matched_keywords=["artificial intelligence"],
        core_matched_keywords=["artificial intelligence"],
        title_keyword_strengths={"artificial intelligence": "core"},
        relevance_score=6,
        relevance_level="中",
    )

    path = export_news(
        [item],
        [item],
        tmp_path / "roc.xlsx",
        export_options=ExportOptions(calendar_mode=CalendarMode.ROC),
    )
    workbook = load_workbook(path)
    all_sheet = workbook["全部新聞"]
    filtered_sheet = workbook["已初步篩選工作表"]

    assert all_sheet["C2"].value.date().isoformat() == "2026-06-08"
    assert "x-roc" in all_sheet["C2"].number_format
    assert filtered_sheet["K3"].value == "artificial intelligence"
    assert filtered_sheet["K3"].fill.fgColor.rgb == "00E6A817"
    assert workbook["篩選設定"]["B7"].value == "roc"
